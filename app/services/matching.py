"""Private resume parsing and explainable skill matching for HireSphere."""
import os
import re
import math
from collections import Counter
from pathlib import Path

from docx import Document
from pypdf import PdfReader


def extract_text_from_file(path):
    """Extract selectable text from a PDF or DOCX resume."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        reader = PdfReader(str(path))
        return "\n".join((page.extract_text() or "") for page in reader.pages).strip()
    if suffix == ".docx":
        document = Document(str(path))
        paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                paragraphs.extend(cell.text for cell in row.cells if cell.text.strip())
        return "\n".join(paragraphs).strip()
    return ""


def get_user_resume_text(user, upload_folder):
    """Read the logged-in user's stored resume without exposing its path publicly."""
    filename = getattr(user, "resume_filename", None)
    if not filename:
        return ""
    safe_name = os.path.basename(filename)
    path = Path(upload_folder) / safe_name
    if not path.is_file():
        return ""
    try:
        return extract_text_from_file(path)
    except Exception:
        return ""


def _skills_list(required_skills):
    # Skills are entered by employers as comma/semicolon/newline-separated values.
    parts = re.split(r"[,;\n|]+", required_skills or "")
    return list(dict.fromkeys(p.strip().lower() for p in parts if p.strip()))


def score_resume(resume_text, required_skills):
    """Return a transparent keyword/skill overlap score and matched/missing skills.

    This is an explainable NLP-style baseline, not a trained predictive model.
    """
    skills = _skills_list(required_skills)
    if not skills:
        return {"score": 0, "matched": [], "missing": []}
    normalized_text = re.sub(r"\s+", " ", (resume_text or "").lower())
    matched, missing = [], []
    for skill in skills:
        # Escape skill names and match boundaries so 'R' does not match 'React'.
        pattern = r"(?<![a-z0-9+#.])" + re.escape(skill) + r"(?![a-z0-9+#.])"
        (matched if re.search(pattern, normalized_text) else missing).append(skill)
    score = round((len(matched) / len(skills)) * 100) if skills else 0
    return {"score": score, "matched": matched, "missing": missing}


def tfidf_similarity(text_a, text_b):
    """Compute cosine similarity between two texts using a small pure-Python TF-IDF model.

    This is an explainable text-similarity signal, not a trained hiring classifier.
    Keeping it dependency-free makes it reliable in the project's Windows environment.
    """
    token_pattern = re.compile(r"[a-zA-Z][a-zA-Z0-9+#.]{1,}")
    tokens_a = token_pattern.findall((text_a or "").lower())
    tokens_b = token_pattern.findall((text_b or "").lower())
    if not tokens_a or not tokens_b:
        return 0.0
    counts_a, counts_b = Counter(tokens_a), Counter(tokens_b)
    vocabulary = set(counts_a) | set(counts_b)
    # With two documents, terms appearing in both receive less weight.
    idf = {term: math.log(1 + 2 / (1 + int(term in counts_a) + int(term in counts_b))) + 1 for term in vocabulary}
    vec_a = {term: (1 + math.log(count)) * idf[term] for term, count in counts_a.items()}
    vec_b = {term: (1 + math.log(count)) * idf[term] for term, count in counts_b.items()}
    dot = sum(value * vec_b.get(term, 0.0) for term, value in vec_a.items())
    norm_a = math.sqrt(sum(value * value for value in vec_a.values()))
    norm_b = math.sqrt(sum(value * value for value in vec_b.values()))
    if not norm_a or not norm_b:
        return 0.0
    return round(max(0.0, min(1.0, dot / (norm_a * norm_b))) * 100, 1)


def recommend_jobs(resume_text, jobs, limit=6):
    results = []
    for job in jobs:
        match = score_resume(resume_text, job.required_skills)
        description = f"{job.title} {job.description} {job.required_skills} {job.company}"
        similarity = tfidf_similarity(resume_text, description)
        # Preserve the transparent skill score and rank with a combined skill/text signal.
        ranking_score = round(match["score"] * 0.7 + similarity * 0.3)
        results.append({"job": job, **match, "similarity": similarity, "ranking_score": ranking_score})
    results.sort(key=lambda item: (item["ranking_score"], item["score"], item["job"].created_at.timestamp() if item["job"].created_at else 0), reverse=True)
    return results[:limit]
